#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera un archivo HTML autocontenido con dos vistas interactivas:
  1) Registro_MP-2026  (columnas B..AQ, filas 8..973)
  2) PMP_2026          (Carta Gantt)
Ambas vistas comparten las mismas funcionalidades: buscador global,
filtros por columna tipo Excel (listas desplegables con casillas),
ordenamiento, contador de filas, leyenda y exportación a CSV.
"""
import json
import os
import sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))

# Origen del libro Excel: 1) argumento CLI  2) copia en el repo  3) ruta de subida
CANDIDATES = [
    sys.argv[1] if len(sys.argv) > 1 else None,
    os.path.join(HERE, "Programacion_MP_2026.xlsm"),
    "/root/.claude/uploads/758a1371-4358-481e-b55d-b82f2a18dc47/73e4b702-Programacio_nMP_2026.xlsm",
]
SRC = next((p for p in CANDIDATES if p and os.path.exists(p)), None)
if not SRC:
    raise SystemExit("No se encontró el archivo Excel. Pase la ruta como argumento.")
OUT = os.path.join(HERE, "programacion_mp_2026.html")

wb = openpyxl.load_workbook(SRC, data_only=True)


def norm(v):
    """Normaliza un valor de celda a texto limpio para el HTML."""
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    if isinstance(v, str):
        return v.strip()
    return str(v)


# ----------------------------------------------------------------------------
# Definición de columnas
# ----------------------------------------------------------------------------
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun",
         "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

# ---- Registro_MP-2026 : columnas B(2) .. AQ(43) ----
reg_cols = []
# Info: B..S  (col 2..19)
reg_info_idx = list(range(2, 20))
for c in reg_info_idx:
    label = norm(wb["Registro_MP-2026"].cell(row=7, column=c).value)
    reg_cols.append({"col": c, "label": label, "group": None, "type": "text"})
# Meses: T..AQ (col 20..43), pares P/R
c = 20
for m in MESES:
    reg_cols.append({"col": c,     "label": "P", "group": m, "type": "status"})
    reg_cols.append({"col": c + 1, "label": "R", "group": m, "type": "status"})
    c += 2

# ---- PMP_2026 : A(1)..AE(31) + AG(33),AH(34),AI(35)  (se omite AF vacía) ----
pmp_cols = []
pmp_info_idx = list(range(1, 20))  # A..S
for c in pmp_info_idx:
    label = norm(wb["PMP_2026"].cell(row=7, column=c).value)
    pmp_cols.append({"col": c, "label": label, "group": None, "type": "text"})
# Meses T..AE (20..31)
c = 20
for m in MESES:
    pmp_cols.append({"col": c, "label": m, "group": None, "type": "status"})
    c += 1
# Resumen AG,AH,AI (33,34,35)
for c in (33, 34, 35):
    label = norm(wb["PMP_2026"].cell(row=7, column=c).value)
    pmp_cols.append({"col": c, "label": label, "group": None, "type": "num"})


def extract(sheet, cols, r0=8, r1=973):
    ws = wb[sheet]
    rows = []
    for r in range(r0, r1 + 1):
        # Solo filas con ID (columna B = 2) no vacío
        if norm(ws.cell(row=r, column=2).value) == "":
            continue
        rows.append([norm(ws.cell(row=r, column=cd["col"]).value) for cd in cols])
    return rows


reg_rows = extract("Registro_MP-2026", reg_cols)
pmp_rows = extract("PMP_2026", pmp_cols)

# Quitamos la clave interna "col" antes de exportar
for cd in reg_cols:
    cd.pop("col", None)
for cd in pmp_cols:
    cd.pop("col", None)

views = {
    "registro": {
        "id": "registro",
        "title": "Registro MP 2026",
        "subtitle": "Registro de mantenciones preventivas realizadas — Hospital Hernán Henríquez Aravena",
        "columns": reg_cols,
        "rows": reg_rows,
        "stickyCols": 2,   # ID + N° Carpeta
        "labelCol": 3,     # índice de "Equipo" para mostrar en chips/filtros (0-based)
    },
    "pmp": {
        "id": "pmp",
        "title": "PMP 2026 · Carta Gantt",
        "subtitle": "Programa de mantención preventiva (carta Gantt) — Hospital Hernán Henríquez Aravena",
        "columns": pmp_cols,
        "rows": pmp_rows,
        "stickyCols": 2,
        "labelCol": 4,
    },
}

data_json = json.dumps(views, ensure_ascii=False)
print("Registro filas:", len(reg_rows), "columnas:", len(reg_cols))
print("PMP filas:", len(pmp_rows), "columnas:", len(pmp_cols))

# ----------------------------------------------------------------------------
# Plantilla HTML  (los datos se inyectan en __DATA__)
# ----------------------------------------------------------------------------
HTML = r"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Programación MP 2026 · HHHA</title>
<style>
  :root{
    --bg:#0f172a; --panel:#ffffff; --ink:#1e293b; --muted:#64748b;
    --line:#e2e8f0; --head:#1e3a5f; --head2:#27496d; --accent:#2563eb;
    --hover:#eff6ff; --chip:#e2e8f0;
  }
  *{box-sizing:border-box}
  body{margin:0;font-family:"Segoe UI",Roboto,Helvetica,Arial,sans-serif;
       color:var(--ink);background:#f1f5f9;font-size:13px}
  header.top{background:linear-gradient(135deg,#1e3a5f,#2563eb);color:#fff;
       padding:14px 20px;box-shadow:0 2px 8px rgba(0,0,0,.15)}
  header.top h1{margin:0;font-size:18px;font-weight:600}
  header.top p{margin:3px 0 0;font-size:12px;opacity:.85}

  .tabs{display:flex;gap:4px;background:#e2e8f0;padding:8px 20px 0}
  .tab{padding:9px 18px;border:none;border-radius:8px 8px 0 0;cursor:pointer;
       background:#cbd5e1;color:#334155;font-size:13px;font-weight:600}
  .tab.active{background:var(--panel);color:var(--accent);box-shadow:0 -2px 4px rgba(0,0,0,.05)}

  .toolbar{display:flex;flex-wrap:wrap;gap:10px;align-items:center;
       padding:10px 20px;background:var(--panel);border-bottom:1px solid var(--line)}
  .toolbar input[type=search]{padding:7px 11px;border:1px solid #cbd5e1;border-radius:8px;
       width:280px;font-size:13px}
  .btn{padding:7px 13px;border:1px solid #cbd5e1;background:#fff;border-radius:8px;
       cursor:pointer;font-size:12px;font-weight:600;color:#334155}
  .btn:hover{background:#f1f5f9}
  .btn.primary{background:var(--accent);color:#fff;border-color:var(--accent)}
  .btn.primary:hover{background:#1d4ed8}
  .count{margin-left:auto;font-size:12px;color:var(--muted);font-weight:600}
  .count b{color:var(--ink)}

  .legend{display:flex;flex-wrap:wrap;gap:7px 14px;padding:8px 20px;background:#f8fafc;
       border-bottom:1px solid var(--line);font-size:11px;color:var(--muted)}
  .legend .lg{display:inline-flex;align-items:center;gap:5px}
  .legend .sw{width:14px;height:14px;border-radius:3px;display:inline-block;border:1px solid rgba(0,0,0,.1)}

  .wrap{padding:0 20px 28px}
  .tablebox{overflow:auto;max-height:calc(100vh - 230px);border:1px solid var(--line);
       border-radius:8px;background:#fff;position:relative}
  table{border-collapse:separate;border-spacing:0;white-space:nowrap}
  th,td{border-right:1px solid var(--line);border-bottom:1px solid var(--line);
       padding:5px 8px;text-align:left;font-weight:400;vertical-align:middle}
  td .cell{max-width:230px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
  td.st,thead th.mes{min-width:30px;max-width:48px;text-align:center}
  thead th{position:sticky;top:0;background:var(--head);color:#fff;font-weight:600;
       z-index:3;font-size:12px}
  thead tr.r2 th{top:29px;background:var(--head2)}
  thead th.grp{text-align:center}
  th .htxt{display:inline-flex;align-items:center;gap:6px}
  .fbtn{border:none;background:rgba(255,255,255,.18);color:#fff;border-radius:4px;
       cursor:pointer;font-size:10px;line-height:1;padding:3px 5px}
  .fbtn:hover{background:rgba(255,255,255,.35)}
  .fbtn.on{background:#fbbf24;color:#1e293b}
  tbody tr:nth-child(even){background:#f8fafc}
  tbody tr:hover{background:var(--hover)}
  tbody td{font-size:12px}
  tbody tr.hidden{display:none}

  /* columnas fijas (sticky-left) */
  th.sticky,td.sticky{position:sticky;left:0;z-index:2}
  thead th.sticky{z-index:5;background:var(--head)}
  td.sticky{background:#fff}
  tbody tr:nth-child(even) td.sticky{background:#f1f5f9}
  tbody tr:hover td.sticky{background:var(--hover)}
  td.idx{color:var(--muted);font-variant-numeric:tabular-nums}

  /* estados / colores */
  td.st{text-align:center;font-weight:700;font-size:11px}
  .st-x   {background:#dbeafe;color:#1e40af}
  .st-r   {background:#dcfce7;color:#166534}
  .st-pm  {background:#ffedd5;color:#9a3412}
  .st-ra  {background:#ede9fe;color:#5b21b6}
  .st-si  {background:#bbf7d0;color:#14532d}
  .st-sira{background:#99f6e4;color:#115e59}
  .st-causa{background:#fee2e2;color:#991b1b}
  .st-fs  {background:#fef9c3;color:#854d0e}
  .st-baja{background:#e2e8f0;color:#475569}
  .st-no  {background:#fecaca;color:#7f1d1d}
  .st-nu  {background:#f1f5f9;color:#64748b}
  .sw-x{background:#dbeafe}.sw-r{background:#dcfce7}.sw-pm{background:#ffedd5}
  .sw-ra{background:#ede9fe}.sw-si{background:#bbf7d0}.sw-sira{background:#99f6e4}
  .sw-causa{background:#fee2e2}.sw-fs{background:#fef9c3}.sw-baja{background:#e2e8f0}
  .sw-no{background:#fecaca}.sw-nu{background:#f1f5f9}

  /* popup filtro */
  .popup{position:absolute;z-index:50;background:#fff;border:1px solid #cbd5e1;
       border-radius:10px;box-shadow:0 10px 30px rgba(0,0,0,.22);width:250px;
       padding:10px;display:none;font-size:12px;color:var(--ink)}
  .popup.show{display:block}
  .popup .srt{display:flex;gap:6px;margin-bottom:8px}
  .popup .srt .btn{flex:1;text-align:center;padding:5px}
  .popup input.fsearch{width:100%;padding:6px 8px;border:1px solid #cbd5e1;
       border-radius:6px;margin-bottom:8px}
  .popup .list{max-height:230px;overflow:auto;border:1px solid var(--line);
       border-radius:6px;padding:4px}
  .popup .opt{display:flex;align-items:center;gap:7px;padding:3px 5px;border-radius:4px;cursor:pointer}
  .popup .opt:hover{background:#f1f5f9}
  .popup .opt input{cursor:pointer}
  .popup .opt span{overflow:hidden;text-overflow:ellipsis}
  .popup .acts{display:flex;gap:6px;margin-top:9px}
  .popup .acts .btn{flex:1;text-align:center}
  .empty{padding:30px;text-align:center;color:var(--muted)}
</style>
</head>
<body>
<header class="top">
  <h1>Programación de Mantención Preventiva 2026</h1>
  <p>Equipos médicos críticos · Hospital Hernán Henríquez Aravena</p>
</header>

<div class="tabs" id="tabs"></div>
<div class="toolbar">
  <input type="search" id="globalSearch" placeholder="🔍 Buscar en todas las columnas...">
  <button class="btn" id="clearFilters">✕ Limpiar filtros</button>
  <button class="btn primary" id="exportCsv">⬇ Exportar CSV</button>
  <span class="count" id="rowCount"></span>
</div>
<div class="legend" id="legend"></div>
<div class="wrap">
  <div class="tablebox" id="tablebox"></div>
</div>

<div class="popup" id="popup"></div>

<script>
const DATA = __DATA__;

/* ---------- utilidades de estado/color ---------- */
function statusClass(v){
  if(!v) return "";
  const t = String(v).trim();
  if(/^C\d$/i.test(t)) return "st-causa";
  const m = {
    "X":"st-x","R":"st-r","PM":"st-pm","RA":"st-ra",
    "Si":"st-si","Sí":"st-si","Si-RA":"st-sira","Sí-RA":"st-sira",
    "FS":"st-fs","Baja":"st-baja","No":"st-no","NU":"st-nu"
  };
  return m[t] || "";
}
const LEGEND = [
  ["sw-x","X – Programado"],["sw-r","R – Reprogramado (P)"],
  ["sw-pm","PM – Pendiente mes"],["sw-ra","RA – Realizado anticipado"],
  ["sw-si","Si – Realizada"],["sw-sira","Si-RA – Realizada (reprog.)"],
  ["sw-causa","C1–C8 – Causa reprogramación"],["sw-fs","FS – Fuera de servicio"],
  ["sw-baja","Baja"],["sw-no","No realizada"],["sw-nu","NU – No utilizado"]
];

/* ---------- estado global ---------- */
let CURRENT = "registro";
const state = {};   // por vista: {filters:{idx:Set}, sort:{idx,dir}, search:""}
for(const k in DATA){ state[k] = {filters:{}, sort:null, search:""}; }

const $ = s => document.querySelector(s);
const tabsEl = $("#tabs"), boxEl = $("#tablebox"), popup = $("#popup");

/* ---------- pestañas ---------- */
function buildTabs(){
  tabsEl.innerHTML = "";
  for(const k in DATA){
    const b = document.createElement("button");
    b.className = "tab" + (k===CURRENT?" active":"");
    b.textContent = DATA[k].title;
    b.onclick = () => { CURRENT = k; buildTabs(); buildLegend(); renderTable(); };
    tabsEl.appendChild(b);
  }
}

function buildLegend(){
  $("#legend").innerHTML = LEGEND.map(([c,t]) =>
    `<span class="lg"><span class="sw ${c}"></span>${t}</span>`).join("");
}

/* ---------- construcción de la tabla ---------- */
function renderTable(){
  const v = DATA[CURRENT];
  const cols = v.columns;
  const hasGroups = cols.some(c => c.group);
  const sticky = v.stickyCols || 0;

  // THEAD
  let thead = "<thead>";
  // fila 1
  thead += "<tr class='r1'>";
  let i = 0;
  while(i < cols.length){
    const c = cols[i];
    if(c.group){
      // agrupar columnas consecutivas con mismo group
      let j = i; while(j < cols.length && cols[j].group === c.group) j++;
      thead += `<th class='grp' colspan='${j-i}'>${esc(c.group)}</th>`;
      i = j;
    }else{
      let cl = (i < sticky ? "sticky " : "") + (c.type==="status" ? "mes" : "");
      thead += `<th class='${cl.trim()}' rowspan='${hasGroups?2:1}' style='${i<sticky?leftStyle(cols,i,sticky):""}'>${headInner(c,i)}</th>`;
      i++;
    }
  }
  thead += "</tr>";
  // fila 2 (solo subcolumnas de grupos)
  if(hasGroups){
    thead += "<tr class='r2'>";
    cols.forEach((c,idx) => { if(c.group) thead += `<th class='mes'>${headInner(c,idx)}</th>`; });
    thead += "</tr>";
  }
  thead += "</thead>";

  // TBODY
  const rowsHtml = v.rows.map((row,ri) => {
    let tds = "";
    cols.forEach((c,idx) => {
      const val = row[idx] ?? "";
      let cls = "", inner;
      if(c.type === "status"){
        const sc = statusClass(val); cls = "st" + (sc?(" "+sc):"");
        inner = esc(val);
      }else{
        inner = `<div class='cell' title="${esc(val)}">${esc(val)}</div>`;
      }
      let style = "";
      if(idx < sticky){ cls += (cls?" ":"") + "sticky"; style = leftStyle(cols,idx,sticky); }
      tds += `<td class='${cls}' style='${style}'>${inner}</td>`;
    });
    return `<tr data-ri='${ri}'>${tds}</tr>`;
  }).join("");

  boxEl.innerHTML = `<table>${thead}<tbody>${rowsHtml}</tbody></table>`;
  applyFilters();
}

function headInner(c, idx){
  const f = state[CURRENT].filters[idx];
  const on = f && f.size ? " on" : "";
  return `<span class='htxt'>${esc(c.label)}`+
         `<button class='fbtn${on}' data-idx='${idx}' title='Filtrar / ordenar'>▼</button></span>`;
}

// desplazamiento acumulado para columnas sticky
function leftStyle(cols, idx, sticky){
  // ancho aproximado por columna previa (se calcula tras render para exactitud)
  return "left:" + (window.__lefts && window.__lefts[idx]!=null ? window.__lefts[idx] : (idx*120)) + "px;";
}

function esc(s){ return String(s).replace(/[&<>"']/g, m =>
  ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[m])); }

/* recalcula offsets exactos de columnas sticky tras render */
function recalcSticky(){
  const v = DATA[CURRENT]; const sticky = v.stickyCols||0;
  if(!sticky) return;
  const firstRow = boxEl.querySelector("tbody tr");
  if(!firstRow) return;
  const tds = firstRow.children;
  let acc = 0; const lefts = [];
  for(let i=0;i<sticky;i++){ lefts[i]=acc; acc += tds[i].offsetWidth; }
  window.__lefts = lefts;   // referencia para futuros renders
  // aplicar offsets exactos a header y cuerpo
  const heads = boxEl.querySelectorAll("thead th.sticky");
  heads.forEach((th,n)=>{ th.style.left = lefts[n] + "px"; });
  boxEl.querySelectorAll("tbody tr").forEach(tr=>{
    for(let i=0;i<sticky;i++){ if(tr.children[i]) tr.children[i].style.left = lefts[i]+"px"; }
  });
}

/* ---------- filtrado ---------- */
function applyFilters(){
  const v = DATA[CURRENT];
  const st = state[CURRENT];
  const search = st.search.trim().toLowerCase();
  const fEntries = Object.entries(st.filters).filter(([,s]) => s && s.size);

  const trs = boxEl.querySelectorAll("tbody tr");
  let shown = 0;
  trs.forEach(tr => {
    const row = v.rows[+tr.dataset.ri];
    let ok = true;
    // filtros por columna
    for(const [idx,set] of fEntries){
      if(!set.has(row[idx] ?? "")){ ok = false; break; }
    }
    // búsqueda global
    if(ok && search){
      ok = row.some(cell => String(cell).toLowerCase().includes(search));
    }
    tr.classList.toggle("hidden", !ok);
    if(ok) shown++;
  });
  $("#rowCount").innerHTML = `Mostrando <b>${shown}</b> de <b>${v.rows.length}</b> equipos`;
  recalcSticky();
}

/* ---------- ordenamiento ---------- */
function sortBy(idx, dir){
  const v = DATA[CURRENT];
  const num = v.columns[idx].type === "num";
  v.rows.sort((a,b)=>{
    let x=a[idx]??"", y=b[idx]??"";
    if(num){ x=parseFloat(x)||0; y=parseFloat(y)||0; return dir==="asc"? x-y : y-x; }
    x=String(x).toLowerCase(); y=String(y).toLowerCase();
    if(x<y) return dir==="asc"?-1:1;
    if(x>y) return dir==="asc"?1:-1;
    return 0;
  });
  state[CURRENT].sort = {idx,dir};
  renderTable();
}

/* ---------- popup de filtro (lista desplegable tipo Excel) ---------- */
let popupIdx = null;
function openPopup(idx, anchor){
  popupIdx = idx;
  const v = DATA[CURRENT];
  const st = state[CURRENT];
  // valores distintos de la columna
  const counts = new Map();
  v.rows.forEach(r => { const val = r[idx] ?? ""; counts.set(val, (counts.get(val)||0)+1); });
  const values = [...counts.keys()].sort((a,b)=>{
    const na=parseFloat(a), nb=parseFloat(b);
    if(!isNaN(na)&&!isNaN(nb)) return na-nb;
    return String(a).localeCompare(String(b),"es");
  });
  const sel = st.filters[idx] ? st.filters[idx] : new Set(values); // sin filtro = todos
  const noFilter = !st.filters[idx];

  let h = `<div class='srt'>
      <button class='btn' data-act='asc'>▲ A-Z</button>
      <button class='btn' data-act='desc'>▼ Z-A</button>
    </div>
    <input class='fsearch' placeholder='Buscar valor...'>
    <div class='list'>
      <label class='opt'><input type='checkbox' class='all' ${ (noFilter||sel.size===values.length)?'checked':'' }><b>(Seleccionar todo)</b></label>`;
  values.forEach(val=>{
    const lbl = val==="" ? "(vacías)" : esc(val);
    const ck = (noFilter || sel.has(val)) ? "checked" : "";
    h += `<label class='opt'><input type='checkbox' class='vchk' data-val='${esc(val)}' ${ck}><span title="${lbl}">${lbl}</span> <span style='color:#94a3b8'>(${counts.get(val)})</span></label>`;
  });
  h += `</div>
    <div class='acts'>
      <button class='btn' data-act='clear'>Limpiar</button>
      <button class='btn primary' data-act='apply'>Aceptar</button>
    </div>`;
  popup.innerHTML = h;
  popup.classList.add("show");

  // posición
  const r = anchor.getBoundingClientRect();
  const px = Math.min(r.left + window.scrollX, window.scrollX + document.documentElement.clientWidth - 270);
  popup.style.left = Math.max(8, px) + "px";
  popup.style.top  = (r.bottom + window.scrollY + 4) + "px";

  // eventos internos
  popup.querySelector(".fsearch").oninput = e=>{
    const q = e.target.value.toLowerCase();
    popup.querySelectorAll(".list .opt").forEach(o=>{
      if(o.querySelector(".all")) return;
      const t = o.textContent.toLowerCase();
      o.style.display = t.includes(q) ? "" : "none";
    });
  };
  popup.querySelector(".all").onchange = e=>{
    popup.querySelectorAll(".vchk").forEach(c=>{ if(c.closest(".opt").style.display!=="none") c.checked = e.target.checked; });
  };
  popup.querySelectorAll("[data-act]").forEach(b=>{
    b.onclick = ()=>{
      const act = b.dataset.act;
      if(act==="asc"||act==="desc"){ closePopup(); sortBy(idx,act); return; }
      if(act==="clear"){
        delete st.filters[idx];
        closePopup(); renderTable(); return;
      }
      if(act==="apply"){
        const chosen = new Set();
        popup.querySelectorAll(".vchk").forEach(c=>{ if(c.checked) chosen.add(c.dataset.val); });
        if(chosen.size===values.length){ delete st.filters[idx]; }
        else { st.filters[idx] = chosen; }
        closePopup(); renderTable(); return;
      }
    };
  });
}
function closePopup(){ popup.classList.remove("show"); popupIdx=null; }

/* ---------- eventos globales ---------- */
boxEl.addEventListener("click", e=>{
  const b = e.target.closest(".fbtn");
  if(b){ e.stopPropagation(); openPopup(+b.dataset.idx, b); }
});
document.addEventListener("click", e=>{
  if(popup.classList.contains("show") && !popup.contains(e.target) && !e.target.closest(".fbtn"))
    closePopup();
});
$("#globalSearch").addEventListener("input", e=>{
  state[CURRENT].search = e.target.value; applyFilters();
});
$("#clearFilters").addEventListener("click", ()=>{
  state[CURRENT].filters = {}; state[CURRENT].search = "";
  $("#globalSearch").value = ""; renderTable();
});
$("#exportCsv").addEventListener("click", exportCsv);
window.addEventListener("resize", recalcSticky);

function exportCsv(){
  const v = DATA[CURRENT];
  const cols = v.columns;
  const head = cols.map(c => (c.group? c.group+" " : "") + c.label);
  const lines = [head];
  boxEl.querySelectorAll("tbody tr").forEach(tr=>{
    if(tr.classList.contains("hidden")) return;
    lines.push(v.rows[+tr.dataset.ri]);
  });
  const csv = lines.map(r => r.map(cell=>{
    const s = String(cell ?? "");
    return /[",\n;]/.test(s) ? '"'+s.replace(/"/g,'""')+'"' : s;
  }).join(";")).join("\r\n");
  const blob = new Blob(["﻿"+csv], {type:"text/csv;charset=utf-8"});
  const a = document.createElement("a");
  a.href = URL.createObjectURL(blob);
  a.download = (CURRENT==="registro"?"Registro_MP_2026":"PMP_2026") + "_filtrado.csv";
  a.click();
}

/* ---------- init ---------- */
buildTabs(); buildLegend(); renderTable();
</script>
</body>
</html>
"""

html_out = HTML.replace("__DATA__", data_json)
with open(OUT, "w", encoding="utf-8") as f:
    f.write(html_out)
print("HTML escrito en:", OUT, "tamaño:", len(html_out), "bytes")
